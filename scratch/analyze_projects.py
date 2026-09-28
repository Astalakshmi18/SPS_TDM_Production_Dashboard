import openpyxl
import pandas as pd
import json

def analyze_fiji():
    print("=== FIJI ===")
    wb = openpyxl.load_workbook(r'media\uploads\gsheet_1KY0YBnI_c178db.xlsx', read_only=True, data_only=True)
    ws = wb['Inventory']
    headers = [c for c in list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]]
    for idx, h in enumerate(headers):
        col_letter = openpyxl.utils.get_column_letter(idx + 1)
        print(f"  {col_letter}: {h}")
    wb.close()

def analyze_smc():
    print("=== SMC ===")
    wb = openpyxl.load_workbook(r'media\uploads\gsheet_1kvJJYAK_f76fc1.xlsx', read_only=True, data_only=True)
    for s in ['Inventory_Index', 'Inventory_ILM', 'Calculation']:
        ws = wb[s]
        headers = [c for c in list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]]
        print(f"--- SMC {s} ---")
        for idx, h in enumerate(headers[:30]):
            col_letter = openpyxl.utils.get_column_letter(idx + 1)
            print(f"  {col_letter}: {h}")
    wb.close()

def analyze_newspaper():
    print("=== NEWSPAPER ===")
    wb = openpyxl.load_workbook(r'media\uploads\gsheet_1FniG5GO_5e8fec.xlsx', read_only=True, data_only=True)
    ws = wb['Inventory']
    headers = [c for c in list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]]
    for idx, h in enumerate(headers):
        col_letter = openpyxl.utils.get_column_letter(idx + 1)
        print(f"  {col_letter}: {h}")
    wb.close()

if __name__ == '__main__':
    analyze_fiji()
    analyze_smc()
    analyze_newspaper()
